"""Build benchmark stimuli from local course sources, without reading answer keys.

Optional preparation dependencies: pymupdf, python-pptx, resvg-py (not runtime dependencies).
Two single-slide PowerPoint PDF exports are retained in data/image_sources.
All crop coordinates are PDF points, with a top-left origin.
"""
from __future__ import annotations

import hashlib
import html
import json
from pathlib import Path
import shutil
import xml.etree.ElementTree as ET

import pymupdf as pdf
import resvg_py
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'data/images'
EDITABLE = ROOT / 'data/image_sources'
RECORDS = []
SCALE = 3  # 216 dpi; preserve fine processor control labels.
NS = '{http://www.w3.org/2000/svg}'
ET.register_namespace('', 'http://www.w3.org/2000/svg')
ET.register_namespace('xlink', 'http://www.w3.org/1999/xlink')


def source(pattern):
    matches = sorted((ROOT / 'referenceFiles').rglob(pattern))
    if len(matches) != 1:
        raise ValueError(f'Expected one source for {pattern}: {matches}')
    return matches[0]


def record(name, paths, location, method):
    RECORDS.append(dict(
        image=f'data/images/{name}.png',
        sources=[dict(path=str(p.relative_to(ROOT)), sha256=hashlib.sha256(p.read_bytes()).hexdigest()) for p in paths],
        source_location=location, preparation=method,
    ))


def save_page(name, page):
    page.get_pixmap(matrix=pdf.Matrix(SCALE, SCALE), alpha=False).save(OUT / f'{name}.png')


def crop_doc(doc, rect, *, masks=(), header=None):
    p = doc[0]
    for area in masks:
        p.add_redact_annot(pdf.Rect(area), fill=(1, 1, 1))
    if masks:
        p.apply_redactions(graphics=2)
    r = pdf.Rect(rect)
    out = pdf.open()
    top = 38 if header else 10
    q = out.new_page(width=r.width + 20, height=r.height + top + 10)
    q.show_pdf_page(pdf.Rect(10, top, 10+r.width, top+r.height), doc, 0, clip=r)
    if header:
        q.insert_text((12, 23), header, fontsize=11, color=(.1,.1,.1))
    return out


def original_page(path, number):
    doc = pdf.open(path)
    single = pdf.open()
    single.insert_pdf(doc, from_page=number-1, to_page=number-1)
    return single


def crop(name, path, number, rect, *, masks=(), header=None, method='Faithful diagram crop; slide heading and teaching prose excluded.'):
    doc = crop_doc(original_page(path, number), rect, masks=masks, header=header)
    save_page(name, doc[0])
    record(name, [path], f'page {number}', method)
    return doc


def svg_document(text):
    svg = pdf.open(stream=text.encode(), filetype='svg')
    return pdf.open(stream=svg.convert_to_pdf(), filetype='pdf')


def vector_clean(path, number, *, remove_blue=False, neutral_fills=False):
    """Remove answer-only vector annotations before rasterizing; retain text labels."""
    p = pdf.open(path)[number-1]
    if remove_blue:
        for block in p.get_text('dict')['blocks']:
            for line in block.get('lines',[]):
                for span in line['spans']:
                    if span['text'].strip() in ('Stall','Trouble!','Flush','these','instructions'):
                        p.add_redact_annot(pdf.Rect(span['bbox']),fill=False)
        p.apply_redactions(graphics=0,images=0)
    # Outlined glyphs preserve exact source fonts and per-character positioning.
    root = ET.fromstring(p.get_svg_image(text_as_path=True))
    for parent in root.iter():
        for node in list(parent):
            tag = node.tag.removeprefix(NS)
            fill, stroke = node.get('fill','').lower(), node.get('stroke','').lower()
            text = ''.join(node.itertext()).strip()
            if remove_blue:
                annotation = (tag == 'path' and (fill in ('#0000ff', '#0070c0') or stroke in ('#0000ff', '#0070c0')))
                annotation |= tag == 'text' and text in ('Stall','Trouble!','Flush','these','instructions')
                # The dotted stall overlay is a tiled SVG pattern, not datapath content.
                annotation |= tag in ('rect','path') and fill.startswith('url(')
                annotation |= tag == 'path' and node.get('d','').startswith('M200.42 26.153H220.67')
                if annotation:
                    parent.remove(node)
                    continue
            if neutral_fills and tag == 'path' and fill not in ('', 'none', '#000000', '#ffffff'):
                node.set('fill','#ffffff')
    return svg_document(ET.tostring(root, encoding='unicode'))


def authored(name, source_paths, location, method):
    path = EDITABLE / f'{name}.svg'
    zoom = 1 if name == 'cs2113_18_architecture_diagram_revised' else SCALE
    (OUT / f'{name}.png').write_bytes(resvg_py.svg_to_bytes(svg_path=str(path),zoom=zoom,background='white'))
    record(name, source_paths + [path], location, method)


def build():
    RECORDS.clear()
    OUT.mkdir(exist_ok=True)
    arith=source('Chapter 4 Arithmetic*')
    micro=source('Chapter 3B*')
    pipe=source('Chapter 5*')
    cache=source('Chapter 7*')
    crop('cg3207_01_adder_symbols',arith,3,(151,111,493,253))
    authored('cg3207_01_ripple_chain_unanswered',[arith],'page 3; turn-4 extension','Unanswered two-column diagram: block types and carry connection left unspecified.')
    crop('cg3207_02_alu_control_circuit',arith,10,(32,126,699,512),masks=[(0,408,350,540)],method='Original ALU and control table retained; explanatory paragraph removed.')
    a=crop('cg3207_03_single_cycle_addi',micro,17,(17,98,656,477))
    b=crop('cg3207_06_jalr_datapath',micro,23,(17,107,656,477))
    comp=pdf.open(); q=comp.new_page(width=680,height=854)
    q.insert_text((12,20),'A',fontsize=16)
    q.show_pdf_page(pdf.Rect(10,27,670,426),a,0)
    q.insert_text((12,448),'B',fontsize=16)
    q.show_pdf_page(pdf.Rect(10,455,670,845),b,0)
    save_page('cg3207_04_processor_a_vs_b',q)
    record('cg3207_04_processor_a_vs_b',[micro],'pages 17 and 23','Full-width vertical A/B comparison; original control widths and ports retained; no instruction capability answers.')
    authored('cg3207_05_beq_trace',[micro],'page 15','Neutral vector redraw of branch comparison and both next-PC routes; no selected path or computed result; decoded byte-offset assumption explicit.')
    crop('cg3207_07_pipelined_registers',pipe,11,(1,94,719,472),masks=[(24,94,86,99)],header='Fetch (F) | Decode (D) | Execute (E) | Memory (M) | Writeback (W)')
    crop('cg3207_08_data_forwarding_circuitry',pipe,24,(1,94,719,539),masks=[(24,94,86,99),(703,523,720,540)],header='Fetch (F) | Decode (D) | Execute (E) | Memory (M) | Writeback (W)')
    for name,n,rect,header in [
        ('cg3207_09_load_use_hazard_timeline',28,(22,321,696,538),None),
        ('cg3207_10_control_hazards_flushing',33,(25,165,693,399),'Branch outcome is determined in Execute.')]:
        clean=vector_clean(pipe,n,remove_blue=True)
        doc=crop_doc(clean,rect,header=header)
        save_page(name,doc[0])
        record(name,[pipe],f'page {n}','Original timeline retained; blue answer arrows, stall shading/label and flush callout removed. Instruction and stage labels retained.')
    for name,n,rect,header in [
        ('cg3207_11_direct_mapping_cache',13,(379,208,695,526),'4096 memory blocks | 128 direct-mapped cache lines | 16 words per block'),
        ('cg3207_12_set_associative_mapping',19,(391,204,695,525),'4096 memory blocks | 64 sets, 2 ways each (128 blocks) | 16 words per block')]:
        clean=vector_clean(cache,n,neutral_fills=True)
        if n == 13:
            # The final prose baseline slightly overlaps the crop's top-left edge.
            clean[0].draw_rect(pdf.Rect(375,205,585,217),color=None,fill=(1,1,1))
        # A wide header provides the givens without a solved mapping or bit count.
        r=pdf.Rect(rect); d=pdf.open(); p=d.new_page(width=680,height=380)
        p.insert_text((16,22),header,fontsize=11)
        p.show_pdf_page(pdf.Rect(160,42,520,366),clean,0,clip=r)
        save_page(name,p)
        record(name,[cache],f'page {n}','Cropped original cache/memory structure; worked mapping prose, solved bit counts and color-coded mapping answers removed; scenario parameters added.')
    page_source=source('CG227Lect8.pptx')
    prepare_page_table(EDITABLE/'cg2271_13_source_slide.pdf',page_source)
    crop('cg2271_14_round_robin_timeline',source('CG2271Lect3.pdf'),21,(26,86,685,289),header='Round robin | Quantum: 2 time units (TU)',method='Original execution blocks, order and durations retained. No arrival times or queue history added.')
    sem_source=source('CG2271Lect6.pptx')
    sem_export=EDITABLE/'cg2271_15_source_slide.pdf'
    doc=crop_doc(original_page(sem_export,1),(36,310,679,496))
    save_page('cg2271_15_semaphore_visualization',doc[0])
    record('cg2271_15_semaphore_visualization',[sem_source,sem_export],'slide 19','Original four state panels, arrows and event captions retained; surrounding explanatory bullets removed.')
    prepare_pwm()
    authored('cg2271_17_uart_receiver_basics',[source('LL-8 UART Programming.pptx')],'slides 4-5','Legible vector redraw of one illustrative UART frame with bit labels and center sample points; optional parity identified; no decoded byte or teaching prose.')
    for name,filename in [('cs2113_18_architecture_diagram','ArchitectureDiagram.png'),('cs2113_19_add_sequence_diagram','add-sequence.png'),('cs2113_20_storage_class_diagram','StorageClassDiagram.png')]:
        p=source(filename); shutil.copyfile(p,OUT/f'{name}.png')
        record(name,[p],'original PNG','Byte-for-byte copy of source diagram; all labels, arrows, legend and activation bars retained.')
    authored('cs2113_18_architecture_diagram_revised',[source('ArchitectureDiagram.png')],'turn-4 variant','Original diagram embedded unchanged; one solid UI-to-File association added as an editable vector overlay.')
    finish()


def prepare_page_table(export_path, src):
    # Geometry from the inspected PowerPoint PDF export of slide 7.
    for revised in (False,True):
        doc=original_page(export_path,1)
        p=doc[0]
        if revised:
            update_page_table(p)
        result=crop_doc(doc,(39,71,696,511),masks=[(20,429,490,461),(20,461,560,500)],header='Page size = frame size = 256 bytes')
        name='cg2271_13_logical_address_translation'+('_revised' if revised else '')
        save_page(name,result[0])
        record(name,[src,export_path],'slide 7'+('; turn 4' if revised else ''),'Original logical memory, page table and physical frames; address formula removed; 256-byte size added.'+(' Page 2 maps to frame 4; physical-memory placement updated consistently.' if revised else ''))


def update_page_table(p):
    # Reuse the source font and baseline; do not change the logical page label.
    courier=next(f for f in p.get_fonts() if 'CourierNewPS-Bold' in f[3])
    fontbuffer=p.parent.extract_font(courier[0])[3]
    p.add_redact_annot(pdf.Rect(375,269,388,290),fill=(217/255,234/255,213/255))
    p.apply_redactions(graphics=0)
    p.insert_font(fontname='sourceCourier',fontbuffer=fontbuffer)
    p.insert_text((376.191,283.44),'4',fontname='sourceCourier',fontsize=16)
    for y,fill in [(120.016,(253/255,229/255,205/255)),(264.016,(1,.8,.6))]:
        p.draw_rect(pdf.Rect(582,y,672,y+48),color=(0,0,0),fill=fill,width=.75)
    p.insert_text((598.125,292.32),'Page 2',fontname='sourceCourier',fontsize=16)


def prepare_pwm():
    src=source('LL-5 PWM Programming.pptx')
    slide=Presentation(src).slides[17]
    picture=slide.shapes[4]
    blob=picture.image.blob
    d=pdf.open(); p=d.new_page(width=780,height=475)
    # Native embedded bitmap is retained, with source ELS labels alongside.
    p.insert_image(pdf.Rect(75,25,635,466),stream=blob,keep_proportion=True)
    p.insert_text((16,93),'CnV',fontsize=14)
    p.draw_line((47,91),(75,91),color=(0,.6,0),width=1)
    p.insert_text((642,249),'ELSnB:ELSnA',fontsize=12)
    p.insert_text((658,268),'= 0b10',fontsize=12)
    p.insert_text((642,396),'ELSnB:ELSnA',fontsize=12)
    p.insert_text((658,415),'= 0b01',fontsize=12)
    save_page('cg2271_16_choosing_pwm_polarity',p)
    record('cg2271_16_choosing_pwm_polarity',[src],'slide 18, embedded diagram','Original embedded counter/waveform diagram extracted; ELS labels retained; teaching prose removed; CnV label supplied explicitly.')


def finish():
    for rec in RECORDS:
        path=ROOT/rec['image']; image=pdf.Pixmap(str(path))
        rec.update(width=image.width,height=image.height,sha256=hashlib.sha256(path.read_bytes()).hexdigest())
    RECORDS.sort(key=lambda r:r['image'])
    (EDITABLE/'manifest.json').write_text(json.dumps({'render_scale':SCALE,'images':RECORDS},indent=2)+'\n')
    scenarios=[json.loads(p.read_text()) for p in sorted((ROOT/'data/scenarios').glob('*.json'))]
    expected={d['initial_image_path'] for d in scenarios}
    expected.update(t['revised_image_path'] for d in scenarios for t in d['turns'] if t.get('revised_image_path'))
    assert expected=={r['image'] for r in RECORDS},'Image manifest must cover exactly the scenario references'
    sections=[]
    for d in scenarios:
        images=[('Initial diagram',d['initial_image_path'])]+[(f"Turn {t['turn_id']} variant",t['revised_image_path']) for t in d['turns'] if t.get('revised_image_path')]
        figs=''.join(f'<figure><figcaption>{html.escape(label)}</figcaption><a href="../{path.removeprefix("data/")}"><img loading="lazy" src="../{path.removeprefix("data/")}" alt="{html.escape(d["title"])} — {html.escape(label)}"></a></figure>' for label,path in images)
        sections.append(f'<section id="{d["scenario_id"]}"><h2>{html.escape(d["scenario_id"])} · {html.escape(d["title"])}</h2>{figs}</section>')
    (EDITABLE/'review.html').write_text('''<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>UREx diagram review</title><style>body{font:16px system-ui;margin:24px auto;padding:0 20px;max-width:1400px;background:#f4f5f7;color:#17202a}section{background:white;padding:20px;margin:24px 0;border:1px solid #d4d8de;border-radius:8px}h1{margin-bottom:8px}h2{font-size:20px}figure{margin:24px 0}figcaption{margin-bottom:12px;color:#555}img{max-width:100%;height:auto;display:block}a{color:#175b94}</style><h1>UREx benchmark diagrams</h1><p>20 scenarios · 23 images. Click an image to inspect it at full resolution. Source provenance and preparation steps are in <a href="manifest.json">manifest.json</a>. This review contains no answer-key text.</p>'''+''.join(sections)+'</html>')
    print(f'Prepared {len(RECORDS)} images; review: data/image_sources/review.html')


if __name__=='__main__':
    build()
